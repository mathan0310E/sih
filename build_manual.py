#!/usr/bin/env python3
"""Build the Naan Mudhalvan Generative AI laboratory record manual (.docx).

The document follows the traditional Indian engineering college practical
record-book layout: thin page border, Times New Roman body text, bordered
index table, page numbers in the footer, and a fixed section structure for
every experiment.
"""
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

SERIF = "Times New Roman"
MONO = "Courier New"

# ---------------------------------------------------------------- low level


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_border(section):
    """Draw the thin page border seen in the reference record book."""
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgBorders")):
        sectPr.remove(old)
    b = _el("w:pgBorders", **{"w:offsetFrom": "page"})
    for edge in ("top", "left", "bottom", "right"):
        b.append(_el("w:" + edge, **{
            "w:val": "single", "w:sz": "6", "w:space": "24",
            "w:color": "000000"}))
    sectPr.append(b)


def footer_page_number(section, first_page_blank=True):
    """Number every page except (optionally) the first page of the section."""
    section.different_first_page_header_footer = first_page_blank
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Page ")
    _style(r, 10)
    r2 = p.add_run()
    _style(r2, 10)
    r2._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = " PAGE "
    r2._r.append(t)
    r2._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    if first_page_blank:
        fp = section.first_page_footer
        fp.is_linked_to_previous = False
        for q in list(fp.paragraphs):
            q.text = ""


def _page_break_before(p):
    """Force the paragraph to start a new page via w:pageBreakBefore."""
    p.paragraph_format.page_break_before = True
    return p


def _consume_break(doc, p):
    """Apply a pending page break to the first paragraph that follows it."""
    if getattr(doc, "_pending_break", False):
        _page_break_before(p)
        doc._pending_break = False
    return p


def _next_paragraph(doc):
    """The last paragraph added to the document body, or None."""
    for child in reversed(doc.element.body):
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            return Paragraph(child, doc)
        if child.tag == qn("w:tbl"):
            return None
    return None


def page_break(doc):
    """Break to a new page.

    Instead of emitting a break-only paragraph (which can overflow at a page
    boundary and produce a blank page), the break is deferred to the next
    paragraph, which is flagged with ``pageBreakBefore``.
    """
    doc._pending_break = True
    return None


def _style(run, size=12, bold=False, italic=False, font=SERIF, color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    return run


def shade(cell, fill="D9D9D9"):
    cell._tc.get_or_add_tcPr().append(_el("w:shd", **{"w:val": "clear", "w:fill": fill}))


def cell_borders(cell, sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = _el("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        borders.append(_el("w:" + edge, **{
            "w:val": "single", "w:sz": sz, "w:space": "0", "w:color": "000000"}))
    tcPr.append(borders)


def para_border(p, sz="4"):
    pPr = p._p.get_or_add_pPr()
    pbdr = _el("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        pbdr.append(_el("w:" + edge, **{
            "w:val": "single", "w:sz": sz, "w:space": "6", "w:color": "000000"}))
    pPr.append(pbdr)


def no_space(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    return p


# ---------------------------------------------------------------- document


def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = SERIF
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:ascii"), SERIF)
    st.element.rPr.rFonts.set(qn("w:hAnsi"), SERIF)
    st.element.rPr.rFonts.set(qn("w:cs"), SERIF)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15

    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(2.6)
    s.right_margin = Cm(2.2)
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.header_distance = Cm(1.25)
    s.footer_distance = Cm(1.25)
    return doc


def add_page(doc):
    """Add a new A4 section (used as an explicit page break)."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(2.6)
    s.right_margin = Cm(2.2)
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.header_distance = Cm(1.25)
    s.footer_distance = Cm(1.25)
    set_border(s)
    footer_page_number(s)
    return s


def P(doc, text="", size=12, bold=False, italic=False, align=None,
      before=0, after=6, line=None, font=SERIF):
    p = doc.add_paragraph()
    _consume_break(doc, p)
    if align is not None:
        p.alignment = align
    no_space(p, before, after, line)
    if text:
        _style(p.add_run(text), size, bold, italic, font)
    return p


def rich(doc, parts, size=12, align=None, before=0, after=6, line=None):
    """parts: list of (text, bold, italic)."""
    p = doc.add_paragraph()
    _consume_break(doc, p)
    if align is not None:
        p.alignment = align
    no_space(p, before, after, line)
    for part in parts:
        text, b, i = (part + (False, False))[:3] if len(part) < 3 else part
        _style(p.add_run(text), size, b, i)
    return p


def bullets(doc, items, size=12, indent=0.7, after=2, style="List Bullet"):
    out = []
    for it in items:
        p = doc.add_paragraph(style=style)
        _consume_break(doc, p)
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.15
        if isinstance(it, str):
            it = [(it, False, False)]
        for text, b, i in it:
            _style(p.add_run(text), size, b, i)
        out.append(p)
    return out


def numbered(doc, items, size=12, indent=0.85, after=2):
    return bullets(doc, items, size, indent, after, style="List Number")


def section_head(doc, text, size=13):
    p = P(doc, text, size=size, bold=True, before=10, after=4)
    return p


def code_block(doc, lines):
    """Monospaced, shaded, bordered code paragraph."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Cm(16.0)
    shade(cell, "F2F2F2")
    cell_borders(cell)
    cell.paragraphs[0].text = ""
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        no_space(p, 0, 0, 1.0)
        p.paragraph_format.left_indent = Cm(0.1)
        _style(p.add_run(ln if ln else " "), 9.0, font=MONO)
    P(doc, "", size=4, after=2)
    return tbl


def data_table(doc, header, rows, widths=None, size=10.5, head_fill="D9D9D9"):
    tbl = doc.add_table(rows=1, cols=len(header))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _fixed_layout(tbl)
    hdr = tbl.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        no_space(p, 1, 1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style(p.add_run(h), size, bold=True)
        shade(hdr[i], head_fill)
        cell_borders(hdr[i])
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            no_space(p, 1, 1)
            if isinstance(val, str) and val.startswith("^"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                val = val[1:]
            _style(p.add_run(str(val)), size)
            cell_borders(cells[i])
    if widths:
        _apply_widths(tbl, widths)
    if len(rows) > 8:
        _repeat_header(tbl)
    return tbl


def _repeat_header(tbl):
    """Repeat the header row when a long table spans more than one page."""
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    for old in trPr.findall(qn("w:tblHeader")):
        trPr.remove(old)
    trPr.append(_el("w:tblHeader", **{"w:val": "true"}))


def _fixed_layout(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    tblPr.append(_el("w:tblLayout", **{"w:type": "fixed"}))


def _apply_widths(tbl, widths_cm):
    """Set column widths on every cell and on the tblGrid (cm -> twips)."""
    twips = [int(round(w * 567)) for w in widths_cm]
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for tw in twips:
        grid.append(_el("w:gridCol", **{"w:w": tw}))
    tbl._tbl.insert(list(tbl._tbl).index(tbl._tbl.tblPr) + 1, grid)
    for row in tbl.rows:
        for i, tw in enumerate(twips):
            if i < len(row.cells):
                row.cells[i].width = Cm(widths_cm[i])


def screenshot_area(doc, caption, height_cm=6.0):
    """Clearly labelled evidence area; never a fabricated screenshot."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Cm(16.0)
    cell_borders(cell)
    p = cell.paragraphs[0]
    no_space(p, 18, 18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(p.add_run("Paste actual experiment screenshot here."), 11, italic=True)
    cap = P(doc, caption, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
            before=2, after=8)
    return tbl, cap


def real_figure(doc, image_path, caption, width_cm=15.0):
    """Place an authentic captured screenshot with a bordered caption."""
    if not os.path.exists(image_path):
        raise FileNotFoundError(image_path)
    p = doc.add_paragraph()
    _consume_break(doc, p)
    no_space(p, 6, 2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(image_path, width=Cm(width_cm))
    cap = P(doc, caption, size=10, italic=True,
            align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=8)
    return cap


def borderless_table(doc, rows, widths, size=11, bold_first_col=False):
    """Two-column layout table with no visible borders (signature blocks)."""
    tbl = doc.add_table(rows=0, cols=len(widths))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _fixed_layout(tbl)
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            q = cells[i].paragraphs[0]
            no_space(q, 6, 6)
            _style(q.add_run(str(val)), size, bold=bold_first_col and i == 0)
    _apply_widths(tbl, widths)
    return tbl
