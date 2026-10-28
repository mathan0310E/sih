"""Export the speaker notes and the audience Q&A as .docx companions."""
import os
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from qa_content import QA, INTRO
import refs

OUT = "/workspace/project/out"
NAVY = RGBColor(0x0F, 0x2D, 0x52)
ORANGE = RGBColor(0xC2, 0x5E, 0x1A)
GREY = RGBColor(0x5F, 0x6B, 0x7A)

TITLES = [
 "Title: Philosophy and Its Role in Education and Society",
 "Abstract and structure of the argument",
 "Why philosophy arises \u2014 and why it persists",
 "What philosophy is",
 "The five classical branches of philosophy",
 "The relationship between philosophy and education",
 "The role of philosophy in education",
 "Five major educational philosophies compared",
 "Classroom applications of each tradition",
 "Philosophy and society: from persons to institutions",
 "Justice, equality, freedom and responsibility",
 "Philosophy in research and knowledge development",
 "Philosophy, publication and research ethics",
 "Examples of ethical and unethical research practice",
 "Philosophy in everyday life",
 "Implications for education and society",
 "Limitations and areas for further study",
 "Conclusion and key takeaways",
 "References \u2014 journal articles and standards",
 "References \u2014 books, and questions for discussion",
]


def setup(doc, title, subtitle):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    h = doc.add_paragraph()
    r = h.add_run(title)
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    doc.add_paragraph()


def notes_doc():
    from pptx import Presentation
    src = os.path.join(OUT, "Philosophy_Education_Society_Kaudilyar_R.pptx")
    prs = Presentation(src)
    doc = Document()
    setup(doc, "Speaker Notes",
          "Philosophy and Its Role in Education and Society \u2014 Kaudilyar R, "
          "Department of Computer Science and Engineering")
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        h = doc.add_paragraph()
        r = h.add_run(f"Slide {i} of {total} \u2014 {TITLES[i-1] if i-1 < len(TITLES) else ''}")
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = ORANGE
        body = slide.notes_slide.notes_text_frame.text.strip()
        for block in body.split("\n"):
            if block.strip():
                p = doc.add_paragraph(block.strip())
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.10
        doc.add_paragraph()
    path = os.path.join(OUT, "Speaker_Notes.docx")
    doc.save(path)
    return path


def qa_doc():
    doc = Document()
    setup(doc, "Anticipated Audience Questions and Answers",
          "Companion to: Philosophy and Its Role in Education and Society \u2014 "
          "Kaudilyar R")
    p = doc.add_paragraph(INTRO)
    p.paragraph_format.space_after = Pt(10)
    p.runs[0].font.size = Pt(10)
    for i, (q, a, src) in enumerate(QA, 1):
        h = doc.add_paragraph()
        r = h.add_run(f"Q{i}.  {q}")
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = NAVY
        h.paragraph_format.space_before = Pt(9)
        h.paragraph_format.space_after = Pt(3)
        pa = doc.add_paragraph(a)
        pa.paragraph_format.space_after = Pt(3)
        pa.paragraph_format.line_spacing = 1.10
        ps = doc.add_paragraph()
        rs = ps.add_run("Sources: " + src)
        rs.font.size = Pt(9)
        rs.font.italic = True
        rs.font.color.rgb = GREY
        ps.paragraph_format.space_after = Pt(6)
    doc.add_page_break()
    h = doc.add_paragraph()
    r = h.add_run("Full reference list (APA 7th edition)")
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY
    for grp, items in [("Journal articles", refs.ARTICLES),
                       ("Books", refs.BOOKS),
                       ("Standards and institutional sources", refs.STANDARDS)]:
        hs = doc.add_paragraph()
        rs = hs.add_run(grp)
        rs.font.size = Pt(11)
        rs.font.bold = True
        rs.font.color.rgb = ORANGE
        for it in items:
            p = doc.add_paragraph(it)
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.size = Pt(9)
    path = os.path.join(OUT, "Audience_QA_and_References.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    print("wrote", notes_doc())
    print("wrote", qa_doc())
